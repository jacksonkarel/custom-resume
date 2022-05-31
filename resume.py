from docx import Document

def resume():
    skills = ['Ontologies and Taxonomies', 'Natural Language Processing/Understanding', 'Chatbots', 'Machine Learning', 'Object Oriented Programming', 'NoSQL', 'SPARQL', 'RDF', 'OWL', 'SQL', 'Python', 'JSON', 'Linux/Unix', 'Shell Scripting', 'Git']
    # document = Document('data/resume/ontology.docx')
    first_column = skills[:7]
    sec_column = skills[8:]
    skills_txt = ""
    for idx in range(len(first_column)):
        first_skill = f"⬝ {first_column[idx]}"
        if idx <= len(sec_column) - 1:
            sec_skill = f"    ⬝ {sec_column[idx]}\n"
            skills_txt = first_skill.join((skills_txt, sec_skill))
        else:
            skills_txt = "".join((skills_txt, first_skill, "\n"))
    print(skills_txt)
    # print(document.tables[0].cell(0,0).text)
    # document.save(output)
